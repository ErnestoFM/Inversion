import os
from pptx import Presentation

pptx_path = r"C:\Users\hatue\Downloads\Inversion\docs\Exposicion_Objetivos_Equipo4_Rediseñada.pptx"
prs = Presentation(pptx_path)

s1 = prs.slides[0]
print("=== SLIDE 1 DETAILS ===")
bg = s1.background
if bg.fill.type:
    print(f"Fill type: {bg.fill.type}")
    if hasattr(bg.fill, 'fore_color') and bg.fill.fore_color:
        print(f"Fill fore_color: {bg.fill.fore_color.rgb}")

for i, shape in enumerate(s1.shapes):
    print(f"Shape {i}: {shape.name} | Type: {shape.shape_type} | Left: {shape.left.inches:.2f}, Top: {shape.top.inches:.2f}, Width: {shape.width.inches:.2f}, Height: {shape.height.inches:.2f}")
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            text = p.text.strip()
            if text:
                fc = p.font.color.rgb if p.font and hasattr(p.font, 'color') and hasattr(p.font.color, 'rgb') else "Default"
                fs = p.font.size.pt if p.font and p.font.size else "Default"
                fn = p.font.name if p.font else "Default"
                print(f"   Txt: '{text[:80]}' | Font: {fn}, Size: {fs}, Color: {fc}")
