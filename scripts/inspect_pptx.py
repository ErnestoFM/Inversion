import os
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

pptx_path = r"C:\Users\hatue\Downloads\Inversion\docs\Exposicion_Objetivos_Equipo4_Rediseñada.pptx"
prs = Presentation(pptx_path)

print(f"Slide Width: {prs.slide_width.inches} inches")
print(f"Slide Height: {prs.slide_height.inches} inches")
print(f"Total Slides: {len(prs.slides)}")

img_count = 0
for idx, slide in enumerate(prs.slides):
    print(f"\n--- SLIDE {idx+1} ---")
    bg = slide.background
    fill = bg.fill
    if fill.type:
        print(f"Background Fill Type: {fill.type}")
        if hasattr(fill, 'fore_color') and fill.fore_color and hasattr(fill.fore_color, 'rgb'):
            print(f"Background Color: {fill.fore_color.rgb}")
            
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                text = p.text.strip()
                if text:
                    font_name = p.font.name if p.font else None
                    font_size = p.font.size.pt if p.font and p.font.size else None
                    font_bold = p.font.bold if p.font else None
                    font_color = p.font.color.rgb if p.font and hasattr(p.font, 'color') and hasattr(p.font.color, 'rgb') else None
                    print(f"  Text: '{text[:80]}' | Font: {font_name} {font_size}pt bold={font_bold} color={font_color}")
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            img_count += 1
            print(f"  [PICTURE Shape] Name: {shape.name}, Width: {shape.width.inches:.2f}, Height: {shape.height.inches:.2f}, Left: {shape.left.inches:.2f}, Top: {shape.top.inches:.2f}")

print(f"\nTotal Images Found: {img_count}")
