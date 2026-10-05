import os
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

pptx_path = r"C:\Users\hatue\Downloads\Inversion\docs\Exposicion_Objetivos_Equipo4_Rediseñada.pptx"
prs = Presentation(pptx_path)

out_img_dir = r"C:\Users\hatue\Downloads\Inversion\docs\brand\extracted_images"
os.makedirs(out_img_dir, exist_ok=True)

for idx, slide in enumerate(prs.slides):
    if idx >= 4:
        break
    print(f"\n==================== SLIDE {idx+1} ====================")
    for shape_idx, shape in enumerate(slide.shapes):
        print(f"Shape {shape_idx}: {shape.name} | Type: {shape.shape_type} | Left: {shape.left.inches:.2f}, Top: {shape.top.inches:.2f}, Width: {shape.width.inches:.2f}, Height: {shape.height.inches:.2f}")
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.strip():
                    print(f"   Paragraph: '{p.text.strip()}'")
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image = shape.image
            image_bytes = image.blob
            ext = image.ext
            filename = f"slide_{idx+1}_img_{shape_idx}.{ext}"
            filepath = os.path.join(out_img_dir, filename)
            with open(filepath, "wb") as f:
                f.write(image_bytes)
            print(f"   [Extracted Image] Saved to: {filepath}")
